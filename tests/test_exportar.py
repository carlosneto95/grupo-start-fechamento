"""Exportação para Excel (Fase 4.4), com o banco de exemplo do conftest.

O que se garante: a planilha é o recorte da tela (filtros, ordenação e escopo),
Despesas leva todas as linhas (não só as desenhadas), texto do Tiny nunca vira
fórmula, os tipos servem para somar e filtrar no Excel, e toda exportação fica
na auditoria.
"""

import io
import sqlite3
from datetime import datetime

import pytest
from openpyxl import load_workbook

from financeiro import dre, exportar, paineis
from financeiro.escopo import SISTEMA
from tests.conftest import _conta, cliente_para, gravar


def _abrir(resposta):
    assert resposta.status_code == 200, resposta.status_code
    assert resposta.mimetype == exportar.MIME_XLSX
    assert resposta.headers["Content-Disposition"].startswith("attachment;")
    return load_workbook(io.BytesIO(resposta.data))


def _detalhe(wb) -> list[dict]:
    """Aba Detalhe como lista de dicts {cabeçalho: valor}."""
    linhas = list(wb["Detalhe"].iter_rows(values_only=True))
    return [dict(zip(linhas[0], l)) for l in linhas[1:]]


@pytest.fixture
def admin(banco_exemplo):
    return cliente_para(banco_exemplo)


TELAS = [
    "/despesas",
    "/receitas/vendas",
    "/receitas/servicos",
    "/dre?ano=2026",
    "/dre/linhas?periodo=01/2026&componente=despesa&bloco=",
    "/admin/usuarios",
]


@pytest.mark.parametrize("url", TELAS)
def test_toda_listagem_exporta_e_mostra_o_botao(admin, url):
    html = admin.get(url).get_data(as_text=True)
    assert 'class="botao-exportar"' in html and "formato=xlsx" in html
    sep = "&" if "?" in url else "?"
    wb = _abrir(admin.get(url + sep + "formato=xlsx"))
    assert wb.sheetnames == ["Resumo", "Detalhe"]


def test_despesas_exporta_todas_as_linhas_e_nao_so_as_da_tela(admin, monkeypatch):
    # A tela corta em LINHAS_NA_TELA; a planilha não.
    monkeypatch.setattr(paineis, "LINHAS_NA_TELA", 3)
    linhas = _detalhe(_abrir(admin.get("/despesas?formato=xlsx")))
    total_tela = paineis.despesas(SISTEMA, _Args())["contas"]
    assert len(linhas) == len(total_tela) > 3


def test_filtros_e_ordenacao_da_url_valem_na_planilha(admin):
    linhas = _detalhe(
        _abrir(admin.get("/despesas?empresa=ALFA&ordenar=valor&direcao=desc&formato=xlsx"))
    )
    assert {l["Empresa"] for l in linhas} == {"ALFA"}
    valores = [l["Valor"] for l in linhas]
    assert valores == sorted(valores, reverse=True)
    resumo = dict(
        _abrir(admin.get("/despesas?empresa=ALFA&formato=xlsx"))["Resumo"].iter_rows(
            values_only=True
        )
    )
    assert resumo["Filtro: empresa"] == "ALFA"


def test_escopo_vale_na_planilha(banco_exemplo):
    c = cliente_para(banco_exemplo, perfil="leitura", empresas=("ALFA",))
    for url in ("/despesas", "/receitas/servicos", "/dre?ano=2026"):
        wb = _abrir(c.get(url + ("&" if "?" in url else "?") + "formato=xlsx"))
        texto = str(list(wb["Detalhe"].iter_rows(values_only=True)))
        assert "BETA" not in texto and "Loja D" not in texto, url
    # Tela de Admin continua de Admin, com ou sem formato=xlsx.
    assert c.get("/admin/usuarios?formato=xlsx").status_code == 403


@pytest.mark.parametrize(
    "fornecedor", ['=HYPERLINK("http://x","clique")', "+1+1", "-2+3", "@SUM(A1)", "\t=1"]
)
def test_texto_do_tiny_nunca_vira_formula(admin, banco_exemplo, fornecedor):
    gravar(
        banco_exemplo, [_conta("ALFA", 90, "COMERCIO-Frete", 1.0, "01/2026", fornecedor=fornecedor)]
    )
    wb = _abrir(admin.get("/despesas?formato=xlsx"))
    celulas = [
        c
        for linha in wb["Detalhe"].iter_rows()
        for c in linha
        if isinstance(c.value, str) and fornecedor in c.value
    ]
    assert len(celulas) == 1
    assert celulas[0].value == "'" + fornecedor and celulas[0].data_type == "s"


def test_tipos_servem_para_somar_e_filtrar_no_excel(admin):
    linhas = _detalhe(_abrir(admin.get("/despesas?formato=xlsx")))
    frete = next(l for l in linhas if l["Fornecedor"] == "Transportes A")
    assert isinstance(frete["Valor"], (int, float)) and frete["Valor"] == 1000
    # Competência fica texto (como data, o Excel a trocaria por 01/MM/AAAA).
    assert frete["Competência"] == "01/2026"
    # Texto comum passa intacto (sem o prefixo de proteção).
    assert not any(str(v).startswith("'") for l in linhas for v in l.values() if v)
    datas = [l["Vencimento"] for l in linhas if l["Vencimento"]]
    assert datas and all(isinstance(d, datetime) for d in datas)


def test_dre_exportada_bate_com_a_tela(admin):
    linhas = _detalhe(_abrir(admin.get("/dre?ano=2026&formato=xlsx")))
    montado = dre.montar(SISTEMA, 2026)
    assert len(linhas) == len(montado["linhas"])
    receita = next(l for l in linhas if l["Linha"] == "Receita")
    esperado = next(l for l in montado["linhas"] if l["def"].chave == "receita")
    assert receita["Acumulado"] == pytest.approx(float(esperado["acumulado"]))
    assert receita["01/2026"] == pytest.approx(float(esperado["valores"][0]))


def test_detalhe_da_dre_exporta_o_total_da_celula(admin):
    wb = _abrir(admin.get("/dre/linhas?periodo=01/2026&componente=receita&bloco=&formato=xlsx"))
    linhas = _detalhe(wb)
    resumo = dict(wb["Resumo"].iter_rows(values_only=True))
    assert resumo["Total"] == pytest.approx(sum(l["Valor"] for l in linhas))


def test_toda_exportacao_vai_para_a_auditoria(admin, banco_exemplo):
    admin.get("/receitas/vendas?empresa=ALFA&formato=xlsx")
    conn = sqlite3.connect(banco_exemplo)
    acao, entidade, novo = conn.execute(
        "SELECT acao, entidade, valor_novo FROM auditoria ORDER BY id DESC LIMIT 1"
    ).fetchone()
    conn.close()
    assert (acao, entidade) == ("exportar", "receitas_vendas")
    assert '"empresa": ["ALFA"]' in novo and '"linhas": 1' in novo


def test_formato_desconhecido_devolve_a_tela(admin):
    r = admin.get("/despesas?formato=pdf")
    assert r.status_code == 200 and r.mimetype == "text/html"


class _Args(dict):
    """MultiDict mínimo, sem filtro nenhum."""

    def getlist(self, _):
        return []

    def to_dict(self, flat=True):
        return {}


def test_diferencas_pos_fechamento_exportam(admin, banco_exemplo):
    admin.post("/fechamento/fechar", data={"competencia": "01/2026", "confirmo": "1"})
    conn = sqlite3.connect(banco_exemplo)
    conn.execute("UPDATE contas_pagar SET valor_centavos = 350000 WHERE id = '4'")
    conn.commit()
    conn.close()
    url = "/fechamento/diferencas?competencia=01/2026"
    assert "formato=xlsx" in admin.get(url).get_data(as_text=True)
    linhas = _detalhe(_abrir(admin.get(url + "&formato=xlsx")))
    assert len(linhas) == 1
    assert (linhas[0]["Valor antes"], linhas[0]["Valor agora"], linhas[0]["Efeito"]) == (
        3000,
        3500,
        500,
    )
