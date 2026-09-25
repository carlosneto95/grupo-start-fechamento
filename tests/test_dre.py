"""DRE gerencial (Fase 4.3), com o banco de exemplo do conftest.

O que se garante aqui: cada coluna mensal é o Dashboard daquele mês (mesmo
separar()), o acumulado é a soma das colunas, e todo número leva a linhas cujo
total é o próprio número — sem isso o drill-down seria só decorativo.
"""

import io

import pytest
from openpyxl import load_workbook
from werkzeug.datastructures import MultiDict

from app import dre, paineis
from app.escopo import SISTEMA, Escopo
from tests.conftest import _conta, cliente_para, gravar


def _linha(d, chave):
    return next(l for l in d["linhas"] if l["def"].chave == chave)


def _drill(escopo, periodo, linha, empresas=None):
    """Total do drill-down de uma linha da DRE, pelo mesmo caminho da tela."""
    args = MultiDict(
        [("periodo", periodo), ("componente", linha.componente), ("bloco", linha.bloco or "")]
        + [("empresa", e) for e in empresas or ()]
    )
    return paineis.dre_linhas(escopo, args)["total"]


def test_meses_vao_ate_o_ultimo_com_receita(banco_exemplo):
    d = dre.montar(SISTEMA, 2026)
    # 02/2026 tem receita pela competência manual da nota 203.
    assert d["meses"] == ["01/2026", "02/2026"]
    assert d["referencia"] == "02/2026" and d["mes_anterior"] == "01/2026"


@pytest.mark.parametrize("mes", ["01/2026", "02/2026"])
def test_coluna_do_mes_bate_com_o_dashboard(banco_exemplo, mes):
    d = dre.montar(SISTEMA, 2026)
    i = d["meses"].index(mes)
    painel = paineis.dashboard(SISTEMA, MultiDict([("competencia", mes)]))
    total = painel["blocos"]["total"]
    for campo in ("receita", "imposto", "despesa", "adm1", "adm2", "resultado"):
        assert _linha(d, campo)["valores"][i] == pytest.approx(total[campo], abs=0.005), campo
    assert d["margem"][i] == pytest.approx(total["margem"], abs=0.005)


def test_resultado_fecha_com_os_componentes(banco_exemplo):
    d = dre.montar(SISTEMA, 2026)
    for i in range(len(d["meses"])):
        v = {
            k: _linha(d, k)["valores"][i]
            for k in ("receita", "imposto", "despesa", "adm1", "adm2", "resultado")
        }
        assert v["resultado"] == pytest.approx(
            v["receita"] - v["imposto"] - v["despesa"] - v["adm1"] - v["adm2"], abs=0.005
        )


def test_acumulado_e_a_soma_dos_meses(banco_exemplo):
    d = dre.montar(SISTEMA, 2026)
    for l in d["linhas"]:
        assert l["acumulado"] == sum(l["valores"])


def test_variacao_mes_contra_mes(banco_exemplo):
    d = dre.montar(SISTEMA, 2026)
    rec = _linha(d, "receita")
    assert rec["var_mm"] == rec["valores"][1] - rec["valores"][0]
    assert rec["var_mm_pct"] == pytest.approx(
        (rec["valores"][1] - rec["valores"][0]) / abs(rec["valores"][0]) * 100
    )


def test_ano_anterior_fora_da_visao_nao_compara(banco_exemplo):
    # 2025 existe no banco (conta 12), mas está abaixo do ANO_MINIMO.
    d = dre.montar(SISTEMA, 2026)
    assert d["mesmo_mes_ano_anterior"] is None
    assert all(l["var_aa"] is None for l in d["linhas"])


@pytest.mark.parametrize("periodo", ["01/2026", "02/2026", "2026"])
def test_total_do_drill_down_e_o_numero_da_celula(banco_exemplo, periodo):
    d = dre.montar(SISTEMA, 2026)
    for l in d["linhas"]:
        # Imposto aponta para a receita (é % dela); resultado vai ao Dashboard.
        if not l["def"].componente or l["def"].campo == "imposto":
            continue
        celula = l["acumulado"] if periodo == "2026" else l["valores"][d["meses"].index(periodo)]
        assert _drill(SISTEMA, periodo, l["def"]) == pytest.approx(celula, abs=0.005), (
            periodo,
            l["def"].chave,
        )


def test_filtro_de_empresa_e_escopo(banco_exemplo):
    so_alfa = dre.montar(SISTEMA, 2026, ["ALFA"])
    escopo_alfa = dre.montar(Escopo(1, "t", "leitura", frozenset({"ALFA"})), 2026)
    assert _linha(so_alfa, "receita")["valores"] == _linha(escopo_alfa, "receita")["valores"]
    # ALFA só tem a venda de 01/2026: 02/2026 some.
    assert so_alfa["meses"] == ["01/2026"]
    assert _linha(so_alfa, "receita_servicos")["acumulado"] == 0


def test_drill_down_rejeita_parametro_invalido(banco_exemplo):
    args = MultiDict([("periodo", "2026"), ("componente", "drop table"), ("bloco", "")])
    assert paineis.dre_linhas(SISTEMA, args)["linhas"] == []


def test_telas_respondem_e_respeitam_escopo(banco_exemplo):
    c = cliente_para(banco_exemplo, perfil="leitura", empresas=("ALFA",))
    r = c.get("/dre?ano=2026")
    assert r.status_code == 200
    assert "BETA" not in r.get_data(as_text=True)  # nem como opção de empresa
    r = c.get("/dre/linhas?periodo=2026&componente=despesa&bloco=servicos")
    assert r.status_code == 200
    assert "Loja D" not in r.get_data(as_text=True)  # conta da BETA
    # Empresa fora do escopo na URL é ignorada, não vaza.
    # (A URL pedida volta no link do botão Excel; o que importa é que BETA não
    # vira opção marcada nem entra nos números — conferido na planilha.)
    r = c.get("/dre?ano=2026&empresa=BETA")
    assert 'value="BETA"' not in r.get_data(as_text=True)
    assert r.get_data(as_text=True).count("checked") == 0
    wb = load_workbook(io.BytesIO(c.get("/dre?ano=2026&empresa=BETA&formato=xlsx").data))
    resumo = dict(wb["Resumo"].iter_rows(values_only=True))
    assert resumo["Empresas"] == "todas do seu acesso"  # BETA descartada
    assert "BETA" not in str(list(wb["Detalhe"].iter_rows(values_only=True)))


def test_mes_de_referencia_corta_colunas_e_move_as_variacoes(banco_exemplo):
    d = dre.montar(SISTEMA, 2026, ate=1)
    assert d["meses"] == ["01/2026"] and d["referencia"] == "01/2026"
    assert d["ultimo_com_receita"] == 2  # o seletor continua oferecendo 02
    assert all(l["var_mm"] is None for l in d["linhas"])  # sem mês anterior
    # O acumulado de jan..ref leva às linhas de jan..ref (lista de períodos).
    desp = _linha(d, "despesa")
    args = MultiDict([("periodo", "01/2026"), ("componente", "despesa"), ("bloco", "")])
    assert paineis.dre_linhas(SISTEMA, args)["total"] == pytest.approx(desp["acumulado"], abs=0.005)


def test_tela_com_mes_de_referencia(banco_exemplo):
    c = cliente_para(banco_exemplo)
    assert c.get("/dre?ano=2026&ate=1").status_code == 200
    assert c.get("/dre?ano=2026&ate=99").status_code == 200  # inválido = padrão
    r = c.get("/dre/linhas?periodo=01/2026&periodo=02/2026&componente=receita&bloco=")
    assert r.status_code == 200 and "01/2026 a 02/2026" in r.get_data(as_text=True)


def test_acumulado_nao_puxa_mes_futuro_sem_receita(banco_exemplo):
    # Parcela lançada em 11/2026 (sem receita): fora das colunas, fora do acumulado
    # e fora do detalhe dele — foi o que a captura com dado real pegou.
    gravar(banco_exemplo, [_conta("BETA", 80, "OBRA X-Material", 12345.0, "11/2026")])
    d = dre.montar(SISTEMA, 2026)
    assert d["meses"] == ["01/2026", "02/2026"]
    c = cliente_para(banco_exemplo)
    html = c.get("/dre?ano=2026").get_data(as_text=True)
    assert "periodo=2026&" not in html  # nenhum link aponta para o ano inteiro
    desp = _linha(d, "despesa_servicos")
    args = MultiDict(
        [("periodo", m) for m in d["meses"]] + [("componente", "despesa"), ("bloco", "servicos")]
    )
    assert paineis.dre_linhas(SISTEMA, args)["total"] == pytest.approx(desp["acumulado"], abs=0.005)
