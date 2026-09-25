"""Exportação para Excel (Fase 4.4): toda listagem vira .xlsx pela própria rota.

Contrato: `?formato=xlsx` na URL de qualquer listagem devolve a planilha com
EXATAMENTE o que a tela mostra — mesmos filtros de cabeçalho, mesma ordenação,
mesmo escopo de empresas —, porque a rota monta o contexto pelo mesmo
`paineis.<tela>()` e só troca o template pela planilha. Não existe consulta
paralela "para exportar" que possa divergir da tela.

Exceção deliberada: Despesas corta a TELA em 2 mil linhas (paineis.
LINHAS_NA_TELA); a planilha leva todas. O corte existe para o HTML não pesar,
e quem exporta quer o recorte inteiro.

Estrutura (padrão do exportar.py do Impostos):
  - aba "Resumo": título, quando e por quem foi gerada, os filtros aplicados
    e os totais da tela;
  - aba "Detalhe": a tabela, com cabeçalho congelado e o autofiltro do Excel.

Tipos de célula: moeda e inteiro viram NÚMERO (dá para somar no Excel);
data "DD/MM/AAAA" do Tiny vira DATA de verdade (dá para filtrar por mês);
competência "MM/AAAA" fica TEXTO — como data, o Excel a trocaria por
01/MM/AAAA e o filtro por "08/2026" deixaria de funcionar.

INJEÇÃO DE FÓRMULA: texto vindo do Tiny (nome de fornecedor, descrição) que
comece com = + - @ (ou tabulação/retorno, que o Excel também aceita como
início de fórmula) seria executado ao abrir o arquivo — um fornecedor
cadastrado como "=HYPERLINK(...)" viraria link malicioso na mão do
contador. Esse texto ganha o prefixo ' e é gravado como texto. É a única
alteração de conteúdo que a planilha faz, e só na célula exportada: o banco
continua espelho do Tiny.

Toda exportação vai para a auditoria (quem, qual tela, quantas linhas e com
quais filtros): exportar é tirar dado do sistema, e a pergunta "quem baixou a
lista de fornecedores?" precisa ter resposta.
"""

from __future__ import annotations

import io
from datetime import date, datetime
from decimal import Decimal

from flask import Response
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app import auditoria
from app.db import FUSO_BRASILIA, get_conn

MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Primeiros caracteres que o Excel/LibreOffice interpretam como fórmula.
PERIGOSOS = ("=", "+", "-", "@", "\t", "\r")

FORMATOS = {
    "moeda": "#,##0.00;[Red]-#,##0.00",
    "pct": "0.0%",
    "data": "DD/MM/YYYY",
    "int": "0",
}
_CABECALHO_FONTE = Font(bold=True, color="FFFFFF")
_CABECALHO_FUNDO = PatternFill("solid", fgColor="1C1B19")


def pedido(args) -> bool:
    """A rota pergunta isto antes de renderizar o template."""
    return args.get("formato") == "xlsx"


def texto_seguro(valor: str) -> str:
    """'=1+1' -> "'=1+1". Texto comum passa intacto."""
    if valor.startswith(PERIGOSOS):
        return "'" + valor
    return valor


def _converter(valor, tipo: str):
    """Valor da linha (já em reais Decimal, datas do Tiny em texto) -> célula."""
    if valor is None or valor == "":
        return None
    if tipo == "moeda":
        # Decimal vai como número exato; o openpyxl grava o texto do número,
        # sem passar por float (nada de 0,1 + 0,2 = 0,30000000000000004).
        return valor if isinstance(valor, Decimal) else Decimal(str(valor))
    if tipo == "pct":
        return Decimal(str(valor)) / 100  # a tela guarda 12,5 (%), o Excel quer 0,125
    if tipo == "int":
        return int(valor)
    if tipo == "data":
        # Tiny: "DD/MM/AAAA"; sistema (último acesso etc.): ISO "AAAA-MM-DD...".
        texto = str(valor)
        try:
            if len(texto) >= 10 and texto[2] == "/" and texto[5] == "/":
                return date(int(texto[6:10]), int(texto[3:5]), int(texto[:2]))
            return date.fromisoformat(texto[:10])
        except ValueError:
            # Data torta do ERP não é consertada aqui: vai como texto, visível.
            return texto_seguro(texto)
    return texto_seguro(str(valor))


def _celula(ws, valor, tipo="texto", negrito=False) -> WriteOnlyCell:
    c = WriteOnlyCell(ws, value=_converter(valor, tipo))
    if isinstance(c.value, str):
        c.data_type = "s"  # nunca fórmula, mesmo que algo escape da regra acima
    if tipo in FORMATOS:
        c.number_format = FORMATOS[tipo]
    if negrito:
        c.font = Font(bold=True)
    return c


def planilha(titulo: str, resumo: list[tuple], colunas: list[tuple], linhas) -> bytes:
    """Monta o .xlsx em memória.

    resumo : [(rótulo, valor, tipo)]           -> aba Resumo
    colunas: [(título, chave_ou_função, tipo)] -> cabeçalho da aba Detalhe;
             a função recebe a linha e devolve o valor (coluna derivada, ex.
             "Considerar" / "Desconsiderar").
    linhas : iterável de dicts

    Modo write_only: grava linha a linha sem montar a planilha inteira em
    memória — Despesas sem filtro passa de 15 mil linhas."""
    wb = Workbook(write_only=True)

    ws = wb.create_sheet("Resumo")
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 60
    t = WriteOnlyCell(ws, value=texto_seguro(titulo))
    t.font = Font(bold=True, size=14)
    ws.append([t])
    ws.append([])
    for rotulo, valor, tipo in resumo:
        ws.append([_celula(ws, rotulo, negrito=True), _celula(ws, valor, tipo)])

    wd = wb.create_sheet("Detalhe")
    for j, (nome, _chave, tipo) in enumerate(colunas, start=1):
        largura = 14 if tipo in FORMATOS else max(12, min(45, len(nome) + 8))
        wd.column_dimensions[get_column_letter(j)].width = largura
    # Congelar e autofiltrar precisam ser declarados antes das linhas no modo
    # write_only. O autofiltro cobre até a última linha (sabida depois).
    wd.freeze_panes = "A2"
    cabecalho = []
    for nome, _chave, _tipo in colunas:
        c = WriteOnlyCell(wd, value=texto_seguro(nome))
        c.font, c.fill = _CABECALHO_FONTE, _CABECALHO_FUNDO
        c.alignment = Alignment(horizontal="center")
        cabecalho.append(c)
    wd.append(cabecalho)
    quantidade = 0
    for item in linhas:
        wd.append(
            [
                _celula(wd, chave(item) if callable(chave) else item.get(chave), tipo)
                for _nome, chave, tipo in colunas
            ]
        )
        quantidade += 1
    wd.auto_filter.ref = f"A1:{get_column_letter(len(colunas))}{quantidade + 1}"

    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()


def _descrever_filtros(filtros: dict) -> list[tuple]:
    """Linhas do Resumo com os filtros de cabeçalho aplicados. Lista longa é
    resumida: o Resumo informa o recorte, a URL completa é que o reproduz."""
    if not filtros:
        return [("Filtros", "nenhum (todas as linhas do seu acesso)", "texto")]
    saida = []
    for coluna, valores in filtros.items():
        texto = ", ".join(valores[:20]) + (
            f" … (+{len(valores) - 20})" if len(valores) > 20 else ""
        )
        saida.append((f"Filtro: {coluna}", texto, "texto"))
    return saida


def enviar(escopo, tela: str, titulo: str, resumo, colunas, linhas, filtros=None) -> Response:
    """Monta a planilha, audita e devolve o download.

    `tela` vira o nome do arquivo e a entidade da auditoria."""
    linhas = list(linhas)
    agora = datetime.now(FUSO_BRASILIA)
    cabeca = [
        ("Gerado em", agora.strftime("%d/%m/%Y %H:%M"), "texto"),
        ("Por", escopo.login, "texto"),
        ("Linhas", len(linhas), "int"),
        *_descrever_filtros(filtros or {}),
        ("", None, "texto"),
    ]
    conteudo = planilha(titulo, cabeca + list(resumo), colunas, linhas)

    conn = get_conn()
    try:
        auditoria.registrar(
            conn,
            "exportar",
            tela,
            None,
            None,
            None,
            {"linhas": len(linhas), "filtros": filtros or {}},
        )
        conn.commit()
    finally:
        conn.close()

    nome = f"{tela}_{agora.strftime('%Y-%m-%d_%H%M')}.xlsx"
    return Response(
        conteudo,
        mimetype=MIME_XLSX,
        headers={
            "Content-Disposition": f'attachment; filename="{nome}"',
            # Planilha com dado financeiro não fica em cache de proxy/navegador.
            "Cache-Control": "no-store",
        },
    )


# ---- colunas de cada tela ----------------------------------------------------------
# Mesmas colunas e mesma ordem da tela; o que a tela mostra como etiqueta ou
# botão vira o texto correspondente. Colunas "originais do Tiny" entram ao lado
# das ajustadas à mão, para a planilha mostrar o que a tela mostra no tooltip.


def _considerar(linha) -> str:
    texto = "Considerar" if linha.get("considerar_efetivo") else "Desconsiderar"
    return texto + (" (manual)" if linha.get("considerar_manual") is not None else "")


COLUNAS_DESPESAS = [
    ("Considerar", _considerar, "texto"),
    ("Empresa", "empresa", "texto"),
    ("Fornecedor", "fornecedor", "texto"),
    ("Emissão", "data_emissao", "data"),
    ("Vencimento", "data_vencimento", "data"),
    ("Liquidação", "data_liquidacao", "data"),
    ("Competência", "competencia", "competencia"),
    ("Valor", "valor", "moeda"),
    ("Categoria", "categoria_primaria", "texto"),
    ("Subcategoria", "subcategoria", "texto"),
    ("Situação", "situacao", "texto"),
]

COLUNAS_NOTAS = [
    ("Considerar", _considerar, "texto"),
    ("Empresa", "empresa", "texto"),
    ("Nº", "numero", "texto"),
    ("Emissão", "data_emissao", "data"),
    ("Cliente", "cliente_nome", "texto"),
    ("Valor", "valor", "moeda"),
    ("Situação", "descricao_situacao", "texto"),
    ("Competência", "competencia_efetiva", "competencia"),
    ("Competência no Tiny", "competencia", "competencia"),
    ("Categoria", "categoria_primaria_efetiva", "texto"),
    ("Categoria no Tiny", "categoria", "texto"),
    ("Marcadores", "marcadores", "texto"),
]

COLUNAS_DRE_LINHAS = [
    ("Empresa", "empresa", "texto"),
    ("Tipo", "tipo", "texto"),
    ("Competência", "competencia", "competencia"),
    ("Nº / vencimento", "documento", "texto"),
    ("Cliente / fornecedor", "descricao", "texto"),
    ("Categoria", "categoria", "texto"),
    ("Valor", "valor", "moeda"),
]


def _lado(qual, campo):
    """Antes/agora de uma diferença pós-fechamento (None = não existia)."""

    def ler(linha):
        lado = linha.get(qual)
        if not lado:
            return None
        if campo == "considerar":
            return "Considerar" if lado.get("considerar") else "Desconsiderar"
        return lado.get(campo)

    return ler


COLUNAS_DIFERENCAS = [
    ("Empresa", "empresa", "texto"),
    ("Tipo", "tipo", "texto"),
    ("Fornecedor / cliente", "descricao", "texto"),
    ("Categoria", "categoria", "texto"),
    ("Mudança", "mudanca", "texto"),
    ("Valor antes", _lado("antes", "valor"), "moeda"),
    ("Considerar antes", _lado("antes", "considerar"), "texto"),
    ("Valor agora", _lado("depois", "valor"), "moeda"),
    ("Considerar agora", _lado("depois", "considerar"), "texto"),
    ("Efeito", "efeito", "moeda"),
]

COLUNAS_USUARIOS = [
    ("Login", "login", "texto"),
    ("Nome", "nome", "texto"),
    ("Perfil", "perfil", "texto"),
    ("Empresas", "empresas_texto", "texto"),
    ("Situação", "situacao", "texto"),
    ("Último acesso", "ultimo_login", "data"),
]


def colunas_dre(meses: list[str]) -> list[tuple]:
    """A DRE é demonstrativo: uma linha por conta, uma coluna por mês."""
    cols = [
        (
            "Linha",
            lambda l: ("    " if l["def"].nivel else "") + l["def"].rotulo.removeprefix("= "),
            "texto",
        )
    ]
    for i, m in enumerate(meses):
        cols.append((m, lambda l, i=i: l["valores"][i], "moeda"))
    cols += [
        ("Acumulado", "acumulado", "moeda"),
        ("Var. m/m", "var_mm", "moeda"),
        ("Var. m/m %", lambda l: None if l["var_mm_pct"] is None else l["var_mm_pct"], "pct"),
        ("Var. a/a", "var_aa", "moeda"),
        ("Var. a/a %", lambda l: None if l["var_aa_pct"] is None else l["var_aa_pct"], "pct"),
    ]
    return cols


COLUNAS_ALERTAS = [
    ("Situação", "situacao", "texto"),
    ("Tipo", "tipo_rotulo", "texto"),
    ("Empresa", "empresa", "texto"),
    ("Fornecedor / categoria", "descricao", "texto"),
    ("Competência", "competencia", "competencia"),
    ("Valor", "valor", "moeda"),
    ("Detalhe", "detalhe", "texto"),
    ("Dispensado por", "dispensado_por", "texto"),
    ("Motivo", "motivo", "texto"),
]
